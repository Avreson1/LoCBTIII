from docx import Document

doc = Document()
doc.add_paragraph('Q: What is the capital of France?')
doc.add_paragraph('A. Berlin')
doc.add_paragraph('B. Paris')
doc.add_paragraph('C. London')
doc.add_paragraph('ANSWER: B')

doc.add_paragraph('Q: What is 2 + 2?')
doc.add_paragraph('A. 3')
doc.add_paragraph('B. 4')
doc.add_paragraph('ANSWER: B')

doc.save('test_upload.docx')
